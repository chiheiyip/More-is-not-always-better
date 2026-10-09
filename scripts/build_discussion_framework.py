"""Render an evidence-backed, locally supplied Discussion framework payload."""
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build(payload, target):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.72)
    sec.left_margin = sec.right_margin = Inches(.78)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        borders = OxmlElement('w:pBdr')
        for edge in ['top', 'bottom', 'left', 'right', 'between']:
            element = OxmlElement('w:' + edge)
            element.set(qn('w:val'), 'nil')
            borders.append(element)
        style.element.get_or_add_pPr().append(borders)
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.1
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 2'].font.size = Pt(11.5)
    for style in ['Heading 1', 'Heading 2']:
        doc.styles[style].paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run(payload.get('footer', 'Discussion framework') + '  |  ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.add_paragraph(payload['title'], 'Title')
    for text in payload['introduction']:
        doc.add_paragraph(text)
    for page in payload['pages']:
        doc.add_page_break()
        doc.add_paragraph(page['heading'], 'Heading 1')
        for block in page['blocks']:
            if 'heading' in block:
                doc.add_paragraph(block['heading'], 'Heading 2')
            for text in block.get('paragraphs', []):
                doc.add_paragraph(text)
    doc.core_properties.title = payload['title']
    doc.core_properties.author = ''
    doc.save(target)


if __name__ == '__main__':
    build(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8-sig')), Path(sys.argv[2]))
