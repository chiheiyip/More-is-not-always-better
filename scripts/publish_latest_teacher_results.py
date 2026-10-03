"""Publish teacher deliverables only from a completed promoted full run.

Research inputs remain in the analysis root. Historical deliverables are kept;
the explicit latest pointer and dated report identify the current package.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
from paper_analysis.teacher.state import file_sha256, git_commit


def dump(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def typed(name, frame):
    return {"name": name, "columns": frame.columns.tolist(),
            "rows": frame.astype(object).where(pd.notna(frame), None).values.tolist()}


def change_status(significance_flip, direction_change):
    # Omnibus F tests have no single coefficient direction. Missing direction
    # is not evidence of a change (bool(float('nan')) would incorrectly be True).
    flags = [False if pd.isna(v) else str(v).lower() in {"true", "1"} for v in (significance_flip, direction_change)]
    return "改变" if any(flags) else "未改变"


def make_docx(markdown: str, target: Path):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.7)
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10 if name == "Normal" else 14 if name == "Title" else 12)
        style.paragraph_format.space_after = Pt(6)
        borders = OxmlElement("w:pBdr")
        for edge in ("top", "left", "bottom", "right", "between"):
            item = OxmlElement(f"w:{edge}")
            item.set(qn("w:val"), "nil")
            borders.append(item)
        style.element.get_or_add_pPr().append(borders)
    document.styles["Normal"].paragraph_format.line_spacing = 1.18
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("|") and index + 1 < len(lines) and "---" in lines[index + 1]:
            values = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(set(cell) <= set("-: ") for cell in row):
                    values.append(row)
                index += 1
            table = document.add_table(rows=len(values), cols=len(values[0]))
            table.style = "Normal Table"
            table.autofit = False
            widths = [4.0, 13.6] if len(values[0]) == 2 else [3.0, 1.3, 2.1, 2.1, 3.4, 3.4, 2.3] if len(values[0]) == 7 else [17.6 / len(values[0])] * len(values[0])
            for c, width in enumerate(widths):
                table.columns[c].width = Cm(width)
            for r, row in enumerate(values):
                for c, text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.width = Cm(widths[c])
                    cell.text = text.replace("`", "")
                    for paragraph in cell.paragraphs:
                        if r == 0:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            try:
                                float(text)
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            except ValueError:
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        paragraph.paragraph_format.space_after = Pt(2)
                        paragraph.paragraph_format.space_before = Pt(2)
                        for run in paragraph.runs:
                            run.font.size = Pt(8.5)
                            run.font.bold = r == 0
                            run.font.color.rgb = RGBColor(32, 44, 58)
                    if r == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E6EDF5")
                        cell._tc.get_or_add_tcPr().append(shade)
                trpr = table.rows[r]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if r == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
            document.add_paragraph()
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], 0)
        elif line.startswith("## "):
            document.add_heading(line[3:], 1)
        elif line:
            document.add_paragraph(line.replace("`", ""))
        index += 1
    document.core_properties.author = ""
    document.core_properties.title = "EEG分母敏感性补充分析"
    document.save(target)


def prepare(root: Path, delivery_root: Path, stamp: str, *, refresh_draft: bool = False):
    summary = json.loads((root / "realdata_run_summary.json").read_text(encoding="utf-8"))
    if summary["status"] != "complete" or not summary["promoted"] or "eeg_denominator_sensitivity" not in summary:
        raise ValueError("A completed promoted full run including denominator sensitivity is required")
    run = Path(summary["run_root"])
    sensitivity = run / "10_eeg_denominator_sensitivity"
    manifest = json.loads((sensitivity / "run_manifest.json").read_text(encoding="utf-8"))
    if manifest["status"] != "complete":
        raise ValueError("Sensitivity analysis is incomplete")
    package = delivery_root / f"{stamp}_老师最新要求_完整更新"
    if package.exists() and (not refresh_draft or (package / "交付文件哈希.json").exists()):
        raise ValueError("Use a new dated package; only unpublished drafts may be refreshed")
    package.mkdir(parents=True, exist_ok=refresh_draft)
    focus = pd.read_csv(sensitivity / "manuscript_comparisons.csv")
    fields = {"表6_额区theta_WWR45×复杂度": "表6：额区θ交互", "枕区theta_复杂度边际效应": "枕区θ：复杂度", "顶区beta_WWR总体效应": "顶区β：WWR总体"}
    table = focus.loc[focus.pair.eq("A→C") & focus.manuscript_claim.isin(fields)].copy()
    rows = []
    for row in table.itertuples():
        pfrom = table.loc[row.Index, "p.value_from"]
        pto = table.loc[row.Index, "p.value_to"]
        rows.append({"EEG结果": fields[row.manuscript_claim], "窗口s": int(row.onset_trim_s),
                     "1–45 β或边际对比": row.estimate_from, "1–40 β或边际对比": row.estimate_to,
                     "1–45 F": row.Fstat_from, "1–40 F": row.Fstat_to,
                     "1–45 p": pfrom, "1–40 p": pto,
                     "1–45 联合q": row.joint_q_from, "1–40 联合q": row.joint_q_to,
                     "结论是否改变": change_status(row.flip_joint_q, row.direction_changed)})
    concise = pd.DataFrame(rows)
    concise.to_csv(package / "老师要求_EEG关键结果对照.csv", index=False, encoding="utf-8-sig")
    claims = pd.read_csv(sensitivity / "manuscript_conclusion_summary.csv")
    claims = claims.loc[claims["比较"].eq("B→C")].reset_index(drop=True)
    facts = pd.DataFrame([
        ("实际band-pass", "0.5–40 Hz"), ("实际notch", "49–51 Hz"),
        ("滤波顺序", "源文件历史存在两种顺序，不能概括为统一顺序；本次未重新滤波。"),
        ("PSD输入", "现有预处理.set/.fdt，已滤波和average reference并运行ICA；未发现ICA成分删除记录。"),
        ("relative分母", "历史A和复现B为1–45 Hz；敏感性C为1–40 Hz。"),
        ("样本", "42人、461共同试次；0/5/10/15 s四窗口；QC与名单固定。"),
        ("factor-level", "等权边际CR2/HTZ重新实现并通过独立参考核对；历史脚本来源缺口保留。"),
        ("previous-scene", "PreviousWWR + PreviousComplexity，保持0805实际执行结构。"),
        ("比较", "B→C检验分母效应；A→B单列源文件复现差异；A→C核对正文。"),
    ], columns=["项目", "实际口径"])
    fractions = pd.read_csv(sensitivity / "power_fraction_summary.csv")
    fractions = fractions.loc[fractions.roi.eq("ALL")]
    comparisons = pd.read_csv(sensitivity / "model_comparisons.csv")
    flip = comparisons.loc[comparisons.pair.eq("B→C") & comparisons.flip_joint_q.eq(True)]
    lines = ["# EEG分母敏感性补充分析", "", f"更新日期：{stamp}。正式全流程：{summary['run_id']}。", "",
             "已按老师要求重新核对relative θ/α/β的1–40 Hz分母结果。主要正文结论的方向和校正显著性保持一致；全部检验中有1项联合q越过0.05，以及1项接近零且不显著的系数方向变化，不能表述为所有统计量完全不变。", "",
             "## 实际处理与统计口径", "", "|项目|实际口径|", "|---|---|"]
    lines.extend(f"|{r[0]}|{r[1]}|" for r in facts.itertuples(index=False, name=None))
    lines += ["", "## 正文关键结果", "", "下表为历史1–45 Hz（A）与1–40 Hz（C）的正文对照；分母效应以完整表中的B→C为准。β适用于单自由度系数或边际对比，顶区β的WWR总体检验报告F。q为四窗口联合BH。", "",
              "|EEG结果|窗口s|1–45 β或F|1–40 β或F|1–45 p / q|1–40 p / q|结论|", "|---|---:|---:|---:|---|---|---|"]
    for _, row in concise.iterrows():
        before = row["1–45 β或边际对比"] if pd.notna(row["1–45 β或边际对比"]) else row["1–45 F"]
        after = row["1–40 β或边际对比"] if pd.notna(row["1–40 β或边际对比"]) else row["1–40 F"]
        lines.append(f"|{row['EEG结果']}|{row['窗口s']}|{before:.5g}|{after:.5g}|{row['1–45 p']:.5g} / {row['1–45 联合q']:.5g}|{row['1–40 p']:.5g} / {row['1–40 联合q']:.5g}|{row['结论是否改变']}|")
    lines += ["", "## 时序和上一场景结论", "",
              "alpha随Block升高、relative theta随Block降低的模式保持。枕区θ复杂度边际效应仍在四窗口为正且显著；顶区β的WWR总体效应仍仅15 s通过扩展联合BH。上一场景48项仍未通过联合BH。", ""]
    for _, row in flip.iterrows():
        lines.append(f"{row.onset_trim_s:g} s额区relative θ的Position效应：β {row.estimate_from:.8g}变为{row.estimate_to:.8g}；p {row['p.value_from']:.8g}变为{row['p.value_to']:.8g}；联合q {row.joint_q_from:.8g}变为{row.joint_q_to:.8g}。窗口内q仍为{row.within_q_to:.8g}，不宜增加跨窗口稳定效应的表述。")
    lines += ["", "5 s的Block1枕区relative θ的WWR45系数接近零，由负变正，p约0.998变为0.997，均不显著。全部配对检验没有原始p或窗口内q的显著性翻转。", "",
              "## 高频功率与分母变化", "", "|窗口s|40–45占总功率均值%|实际分母减少均值%|relative增加均值%|", "|---:|---:|---:|---:|"]
    lines.extend(f"|{r.onset_trim_s:g}|{r.percent_40_45_mean:.6f}|{r.denominator_reduction_percent_mean:.6f}|{r.relative_increase_percent_mean:.6f}|" for r in fractions.itertuples())
    lines += ["", "保留原频率掩码和trapz。40–45 Hz功率占比与分母减少比例分别积分；频率点边界之间的梯形面积使两者不同，未统一按0.15%缩放。", "",
              "## 验证和结果来源", "", "A版本重现240条核心Model1系数，factor-level通过528条独立参考核对。B/C样本、协变量、分子和absolute数据一致。旧结果只有输入、统计方法和文件哈希通过核对后才复用。", "",
              "完整统计结果、置信区间、自由度、p/q及复用证明均随结果保存。正式运行和已验证敏感性模型分别保留计算版本，见CODE_VERSION.json。"]
    brief = "\n".join(lines) + "\n"
    (package / "老师最新要求_EEG分母敏感性.md").write_text(brief, encoding="utf-8")
    make_docx(brief, package / "老师最新要求_EEG分母敏感性.docx")
    for source_name, destination_name in [
        ("论文数据分析结果报告.md", f"论文数据分析结果报告-{stamp}.md"),
        ("老师任务完成矩阵.xlsx", f"老师任务完成矩阵-{stamp}.xlsx"),
        ("realdata_run_summary.json", "正式全流程运行摘要.json"),
    ]:
        shutil.copyfile(root / source_name, package / destination_name)
    sensitivity_files = ["EEG_1-40Hz_完整对照.xlsx", "model_comparisons.csv", "manuscript_comparisons.csv",
                         "manuscript_conclusion_summary.csv", "comparison_summary.csv", "power_fraction_summary.csv",
                         "input_reproduction_differences.csv", "historical_coefficient_regression.csv",
                         "factor_independent_regression.csv", "reuse_verification.json", "verification_summary.json"]
    for name in sensitivity_files:
        if (sensitivity / name).exists():
            shutil.copyfile(sensitivity / name, package / name)
    for version in "ABC":
        target = package / version
        target.mkdir(exist_ok=refresh_draft)
        for name in ["coefficients.csv", "factor_tests.csv", "contrast_matrices.csv", "diagnostics.csv", "family_status.csv"]:
            shutil.copyfile(sensitivity / version / name, target / name)
    dump(package / "CODE_VERSION.json", {"publication_git_sha": git_commit(REPO), "formal_run_git_sha": summary["git_commit"],
                                        "sensitivity_analysis_git_sha": manifest["git_sha"], "run_id": summary["run_id"]})
    dump(package / "teacher_workbook_tables.json", {"tables": [typed("正文关键结果", concise), typed("结论变化", claims), typed("预处理及说明", facts)]})
    dump(package / "publication_state.json", {"outputs_root": str(root), "run_root": str(run), "sensitivity_root": str(sensitivity), "stamp": stamp})
    print(str(package))
    return package


def finalize(package: Path, delivery_root: Path):
    state = json.loads((package / "publication_state.json").read_text(encoding="utf-8"))
    root, run, stamp = Path(state["outputs_root"]), Path(state["run_root"]), state["stamp"]
    required = ["老师最新要求_EEG分母敏感性.docx", "老师最新要求_EEG结果对照.xlsx", "EEG_1-40Hz_完整对照.xlsx", f"结果文件总索引-{stamp}.xlsx"]
    if not all((package / name).is_file() and (package / name).stat().st_size > 0 for name in required):
        raise ValueError("Required rendered report and workbooks must exist before publishing")
    index = pd.read_csv(package / f"结果文件总索引-{stamp}.csv")
    if not all(file_sha256(package / r.文件) == r.SHA256 for r in index.itertuples()):
        raise ValueError("A deliverable changed after indexing")
    # Root-level current entries are exact copies from the finalized package.
    current = [f"论文数据分析结果报告-{stamp}.md", f"老师任务完成矩阵-{stamp}.xlsx", f"结果文件总索引-{stamp}.xlsx",
               "老师最新要求_EEG分母敏感性.docx", "老师最新要求_EEG分母敏感性.md", "老师最新要求_EEG结果对照.xlsx", "CODE_VERSION.json"]
    for name in current:
        shutil.copyfile(package / name, delivery_root / name)
    records = [{"path": p.relative_to(package).as_posix(), "sha256": file_sha256(p), "size_bytes": p.stat().st_size}
               for p in sorted(package.rglob("*")) if p.is_file() and not p.name.endswith((".zip", "tables.json", ".png"))]
    dump(package / "交付文件哈希.json", records)
    archive = package / f"老师最新要求_完整交付包-{stamp}.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as handle:
        for record in records:
            handle.write(package / record["path"], record["path"])
        handle.write(package / "交付文件哈希.json", "交付文件哈希.json")
    with zipfile.ZipFile(archive) as handle:
        if handle.testzip() is not None:
            raise ValueError("Delivery ZIP integrity check failed")
        for record in records:
            if hashlib.sha256(handle.read(record["path"])).hexdigest() != record["sha256"]:
                raise ValueError("Delivery ZIP hash mismatch")
    pointer = {"status": "complete", "package": str(package), "zip": str(archive), "zip_sha256": file_sha256(archive),
               "source_run": str(run), "source_sensitivity": state["sensitivity_root"], "publication_git_sha": git_commit(REPO)}
    dump(delivery_root / "latest_delivery.json", pointer)
    (delivery_root / "README_最新交付.md").write_text(
        f"# 当前老师交付\n\n本次交付：`{package.name}`。先读 `老师最新要求_EEG分母敏感性.docx` 和 `老师最新要求_EEG结果对照.xlsx`。\n\n"
        f"完整多模态报告为 `论文数据分析结果报告-{stamp}.md`，来源是已验收并更新到正式结果目录的 `{run.name}`。\n\n"
        "0805、0809和20261002目录保留作历史追溯，不作为本次最新入口。\n", encoding="utf-8")
    shutil.copyfile(package / "老师最新要求_EEG分母敏感性.md", root / "老师最新要求_EEG分母敏感性.md")
    summary = json.loads((root / "realdata_run_summary.json").read_text(encoding="utf-8"))
    summary["teacher_delivery"] = pointer
    dump(root / "realdata_run_summary.json", summary)
    pd.DataFrame([{ "run_id": summary["run_id"], "status": summary["status"], "git_commit": summary["git_commit"],
                   "questionnaire_participants": summary["questionnaire"]["formal_participants"],
                   "eye_participants": summary["eye"]["primary_participants"], "eeg_participants": summary["eeg"]["scene_qc_participants"],
                   "sensitivity_participants": 42, "sensitivity_common_trials": 461}]).to_csv(root / "realdata_run_summary.csv", index=False, encoding="utf-8-sig")
    sensitivity_summary = json.loads((Path(state["sensitivity_root"]) / "summary.json").read_text(encoding="utf-8"))
    dump(root / "eeg_denominator_sensitivity_latest.json", {"status": "complete", "run_root": state["sensitivity_root"],
                                                          "formal_run_root": str(run), "delivery_root": str(package), "delivery_zip": str(archive),
                                                          "major_manuscript_conclusions_changed": sensitivity_summary["major_manuscript_conclusions_changed"],
                                                          "joint_q_flips_B_to_C": sensitivity_summary["joint_q_flips_B_to_C"], "scope": "integrated promoted full results"})
    print(json.dumps(pointer, ensure_ascii=False))


def build_index(package: Path, pr: str | None):
    state = json.loads((package / "publication_state.json").read_text(encoding="utf-8"))
    stamp = state["stamp"]
    version = json.loads((package / "CODE_VERSION.json").read_text(encoding="utf-8"))
    version["repository_git_sha"] = git_commit(REPO)
    if pr:
        version["pr"] = pr
    dump(package / "CODE_VERSION.json", version)
    rows = []
    for file in sorted(package.rglob("*")):
        if not file.is_file() or file.name.endswith((".png", ".zip", "tables.json", ".ndjson")):
            continue
        if file.name.startswith("结果文件总索引") or file.name in {"交付文件哈希.json", "publication_state.json"}:
            continue
        rows.append({"文件": file.relative_to(package).as_posix(), "来源版本": state["run_root"].split("\\")[-1],
                     "大小字节": file.stat().st_size, "SHA256": file_sha256(file)})
    frame = pd.DataFrame(rows)
    frame.to_csv(package / f"结果文件总索引-{stamp}.csv", index=False, encoding="utf-8-sig")
    dump(package / "delivery_index_tables.json", {"tables": [typed("交付文件", frame)]})
    print("indexed_deliverables", len(frame))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("phase", choices=["prepare", "index", "finalize"])
    parser.add_argument("--outputs-root", type=Path)
    parser.add_argument("--delivery-root", required=True, type=Path)
    parser.add_argument("--stamp", default="20261003")
    parser.add_argument("--package", type=Path)
    parser.add_argument("--refresh-draft", action="store_true")
    parser.add_argument("--pr")
    options = parser.parse_args()
    if options.phase == "prepare":
        prepare(options.outputs_root, options.delivery_root, options.stamp, refresh_draft=options.refresh_draft)
    elif options.phase == "index":
        build_index(options.package, options.pr)
    else:
        finalize(options.package, options.delivery_root)
