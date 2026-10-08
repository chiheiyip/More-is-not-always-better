"""Verify workbook values, source copies and an optional published ZIP."""
import argparse
import hashlib
import json
import math
from pathlib import Path
import sys
import zipfile

import openpyxl

REPO=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(REPO/'src'))
from paper_analysis.teacher.fresh import dump,git_sha
from paper_analysis.teacher.state import file_sha256


def verify(run, delivery=None):
    staged=delivery or run/'handoff_staging'
    payload=json.loads((run/'handoff_tables.json').read_text(encoding='utf-8'))
    sources=json.loads((run/'source_manifest.json').read_text(encoding='utf-8'))
    sources_by_id={row['Source ID']:row for row in sources}
    checks=[]
    for path in [staged/'数据来源交接索引.xlsx', run/'结果文件总索引.xlsx']:
        book=openpyxl.load_workbook(path,data_only=False)
        assert book.sheetnames==[t['name'] for t in payload['tables']]
        cells=links=0
        for table in payload['tables']:
            sheet=book[table['name']]
            assert [c.value for c in sheet[1]]==table['columns']
            assert sheet.max_row==len(table['rows'])+1
            assert sheet.freeze_panes=='A2' and len(sheet.tables)==1
            for r,row in enumerate(table['rows'],2):
                for c,expected in enumerate(row,1):
                    actual=sheet.cell(r,c).value
                    assert actual not in ['#REF!','#VALUE!','#DIV/0!','#NAME?','#NUM!','#N/A']
                    column=table['columns'][c-1]
                    if column in ['flat_name','平铺文件']:
                        assert isinstance(actual,str) and actual.startswith('=HYPERLINK(') and str(expected) in actual
                        assert (staged/'filesource_flat'/str(expected)).is_file()
                        if path.name=='数据来源交接索引.xlsx':assert '"filesource_flat/' in actual
                        links+=1
                    elif isinstance(expected,(int,float)) and not isinstance(expected,bool):
                        assert isinstance(actual,(int,float)) and math.isclose(actual,expected,rel_tol=1e-12,abs_tol=1e-12),(table['name'],r,c,actual,expected)
                    else:
                        assert actual==expected or (expected=='' and actual is None),(table['name'],r,c,actual,expected)
                    cells+=1
        checks.append({'path':str(path),'sha256':file_sha256(path),'cells':cells,'links':links,'status':'passed'})
    for row in sources:
        assert file_sha256(row['source_path'])==row['SHA256']
        assert file_sha256(staged/'filesource_flat'/row['flat_name'])==row['SHA256']
    mapping=next(t for t in payload['tables'] if t['name']=='段落来源映射')
    source_col=mapping['columns'].index('Source ID')
    assert all(row[source_col] in sources_by_id for row in mapping['rows'])
    assert file_sha256(run/'论文数据分析结果报告.md')==file_sha256(staged/'论文数据分析结果报告.md')
    if delivery:
        expected={'论文数据分析结果报告.md','数据来源交接索引.xlsx','数据来源交接说明.docx','filesource_flat'}
        zips=list(delivery.glob('*.zip'));assert len(zips)==1
        assert {p.name for p in delivery.iterdir()}==expected|{zips[0].name}
        with zipfile.ZipFile(zips[0]) as archive:
            assert archive.testzip() is None
            filenames={p.relative_to(delivery).as_posix() for p in delivery.rglob('*') if p.is_file() and p.suffix!='.zip'}
            assert set(archive.namelist())==filenames
            for name in filenames:
                assert hashlib.sha256(archive.read(name)).hexdigest()==file_sha256(delivery/name)
        checks.append({'path':str(zips[0]),'sha256':file_sha256(zips[0]),'files':len(filenames),'status':'passed'})
    dump(run/('published_data_verification.json' if delivery else 'artifact_data_verification.json'),
         {'status':'passed','verification_sha':git_sha(REPO),'sources':len(sources),'workbooks_and_archive':checks})
    print(json.dumps({'status':'passed','sources':len(sources),'checks':len(checks)}))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--run-root',type=Path,required=True)
    parser.add_argument('--delivery-root',type=Path)
    args=parser.parse_args();verify(args.run_root,args.delivery_root)
