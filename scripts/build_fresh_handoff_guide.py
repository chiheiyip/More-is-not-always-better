"""Build the compact teacher reading guide with the bundled document runtime."""
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def main():
    source, target = map(Path, sys.argv[1:])
    guide = json.loads(source.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ["Normal", "Title", "Heading 1"]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        borders = OxmlElement("w:pBdr")
        for edge in ["top", "bottom", "left", "right", "between"]:
            item = OxmlElement("w:" + edge)
            item.set(qn("w:val"), "nil")
            borders.append(item)
        style.element.get_or_add_pPr().append(borders)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.12
    doc.add_paragraph(guide[0], "Title")
    for text in guide[1:]:
        doc.add_paragraph(text)
    doc.core_properties.title = guide[0]
    doc.core_properties.author = ""
    doc.save(target)


if __name__ == "__main__":
    main()
