import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pytest

from paper_analysis.teacher.hf_qc_delivery import validate_numeric_claims
from paper_analysis.teacher.manuscript_document import build, inspect
from paper_analysis.teacher.state import StageBlockedError


def test_numeric_claims_require_exact_source_row_and_full_precision(tmp_path):
    source = tmp_path / "中文.csv"
    pd.DataFrame([{"window": 10, "outcome": "枕区θ", "estimate": .004123456789},
                  {"window": 0, "outcome": "枕区θ", "estimate": .003}]).to_csv(source, index=False, encoding="utf-8-sig")
    sources = {"theta": {"source_path": str(source), "Source ID": "SRC0001", "flat_name": "SRC0001.csv"}}
    claim = {"source": "theta", "filters": {"window": 10, "outcome": "枕区θ"}, "field": "estimate",
             "expected": .004123456789, "location": "3.3 paragraph 2"}
    result = validate_numeric_claims({"numeric_claims": [claim]}, sources)
    assert result.iloc[0]["核验"] == "一致"
    with pytest.raises(StageBlockedError, match="differs"):
        validate_numeric_claims({"numeric_claims": [dict(claim, expected=.0042)]}, sources)
    with pytest.raises(StageBlockedError, match="one source row"):
        validate_numeric_claims({"numeric_claims": [dict(claim, filters={"outcome": "枕区θ"})]}, sources)


def test_appendix_table_insertion_preserves_existing_table_and_media(tmp_path):
    original = tmp_path / "original.docx"
    xml = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
    for text in ["Table B.2", "Note. Earlier result", "Data availability"]:
        xml += f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p>'
    xml += '<w:tbl><w:tblGrid><w:gridCol w:w="9000"/></w:tblGrid>'
    for text in ["Header", "Original value"]:
        xml += f'<w:tr><w:tc><w:tcPr><w:tcW w:w="9000" w:type="dxa"/></w:tcPr><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc></w:tr>'
    xml += '</w:tbl></w:body></w:document>'
    with ZipFile(original, "w") as z:
        z.writestr("word/document.xml", xml)
        z.writestr("word/media/image.png", b"original figure")
    before = original.read_bytes()
    rows = [["Window", "RMS", "PTP"], ["0", "10.0001", "90.0002"], ["5", "9.9", "89.8"]]
    payload = {"source_sha256": hashlib.sha256(before).hexdigest(), "table_insertions": [{
        "before_index": 2, "expected_anchor": "Data availability", "table_template_index": 0,
        "caption_style_index": 0, "note_style_index": 1, "caption": "Table B.3 QC thresholds", "note": "Note. No-HF.", "rows": rows}]}
    target = tmp_path / "revised.docx"
    build(original, payload, target)
    result = inspect(target)
    assert original.read_bytes() == before
    assert result["tables"] == [rows, [["Header"], ["Original value"]]]
    assert result["paragraphs"] == ["Table B.2", "Note. Earlier result", "Table B.3 QC thresholds", "Note. No-HF.", "Data availability"]
    with ZipFile(target) as z: assert z.read("word/media/image.png") == b"original figure"


def test_empty_numerical_evidence_cannot_prepare_a_scientific_package():
    with pytest.raises(StageBlockedError, match="No source-backed"):
        validate_numeric_claims({"numeric_claims": []}, {})


def test_writing_run_is_separate_and_legacy_config_is_compatible(tmp_path):
    from paper_analysis.teacher.hf_qc_delivery import scientific_run
    source, writing = tmp_path / 'science', tmp_path / 'writing'
    assert scientific_run({'run_root': str(source)}) == source
    assert scientific_run({'run_root': str(writing), 'scientific_run_root': str(source)}) == source
    for a, b in [(source, source/'nested'), (source/'nested', source)]:
        with pytest.raises(StageBlockedError, match='disjoint'):
            scientific_run({'run_root': str(a), 'scientific_run_root': str(b)})


def test_later_delivery_reads_proofs_from_scientific_source(tmp_path, monkeypatch):
    from paper_analysis.teacher import hf_qc_delivery as module
    source, writing = tmp_path/'science', tmp_path/'writing'
    (source/'complete_verification').mkdir(parents=True)
    (source/'bootstrap_repeat').mkdir()
    outputs = tmp_path/'outputs'; outputs.mkdir()
    (outputs/'realdata_run_summary.json').write_text('{}')
    (source/'run_manifest.json').write_text(json.dumps({'status':'analysis_complete_pending_delivery'}))
    (source/'methodological_gate.json').write_text(json.dumps({'adopt_NoHF':True}))
    proof = {'status':'passed', 'scopes':{k:{'status':'passed'} for k in ['eye','eeg','joint']},
             'evidence':[], 'eye_reuse':[], 'checks':[]}
    proof_path = source/'complete_verification/verification.json'
    proof_path.write_text(json.dumps(proof))
    manuscript = tmp_path/'source.docx'; manuscript.write_bytes(b'protected')
    payload = tmp_path/'payload.json'
    payload.write_text(json.dumps({'manuscript':{'source_sha256':hashlib.sha256(manuscript.read_bytes()).hexdigest()},'sources':[]}))
    config = {'run_root':str(writing),'scientific_run_root':str(source),'outputs_root':str(outputs),
              'work_root':str(tmp_path/'work'),'delivery_root':str(tmp_path/'delivery'),
              'archive_path':str(tmp_path/'archive'),'writing_payload':str(payload),
              'writing_payload_sha256':hashlib.sha256(payload.read_bytes()).hexdigest(),'manuscript_file':str(manuscript)}
    monkeypatch.setattr(module,'verify_seal',lambda _: None)
    calls=[]
    monkeypatch.setattr(module,'verify_complete_bootstrap_repeat',lambda path:calls.append(path))
    monkeypatch.setattr(module,'validate_numeric_claims',lambda *args:None)
    before = {p:p.read_bytes() for p in source.rglob('*') if p.is_file()}
    assert module.check(config) == []
    assert calls == [source/'bootstrap_repeat/repeat_execution.json']
    assert not writing.exists()
    assert all(p.read_bytes()==value for p,value in before.items())
    proof_path.write_text(json.dumps(dict(proof,status='failed')))
    with pytest.raises(StageBlockedError,match='must all pass'): module.check(config)


@pytest.mark.parametrize('fault', [None, 'missing', 'duplicate', 'header', 'note'])
def test_caption_located_table_reorder_moves_values_and_note_together(tmp_path, fault):
    from docx import Document
    doc=Document()
    doc.add_paragraph('Table B.2. Previous scene')
    table=doc.add_table(rows=2, cols=1); table.cell(0,0).text='Coefficient'; table.cell(1,0).text='0.123'
    doc.add_paragraph('Note. Previous scene retained.')
    doc.add_paragraph('Table B.3. QC thresholds')
    table=doc.add_table(rows=2, cols=1); table.cell(0,0).text='RMS'; table.cell(1,0).text='16.286598128'
    doc.add_paragraph('Note. Frozen NoHF.')
    if fault=='duplicate': doc.add_paragraph('Table B.3. QC thresholds')
    original, target=tmp_path/'original.docx', tmp_path/'revised.docx'; doc.save(original)
    before=original.read_bytes()
    payload={'source_sha256':hashlib.sha256(before).hexdigest(), 'table_reorders':[{
        'caption':'absent' if fault=='missing' else 'Table B.3. QC thresholds',
        'before_caption':'Table B.2. Previous scene',
        'expected_header':['Wrong' if fault=='header' else 'RMS'],
        'expected_note':'Wrong' if fault=='note' else 'Note. Frozen NoHF.'}],
        'paragraph_edits':[{'match_text':True,'expected':'Table B.3. QC thresholds','replacement':'Table B.2. QC thresholds'},
                           {'match_text':True,'expected':'Table B.2. Previous scene','replacement':'Table B.3. Previous scene'}]}
    if fault:
        with pytest.raises(ValueError): build(original,payload,target)
        assert not target.exists()
    else:
        build(original,payload,target)
        result=inspect(target)
        assert result['tables']==[[['RMS'],['16.286598128']],[['Coefficient'],['0.123']]]
        assert result['paragraphs']==['Table B.2. QC thresholds','Note. Frozen NoHF.',
                                       'Table B.3. Previous scene','Note. Previous scene retained.']
    assert original.read_bytes()==before


def test_compact_final_bibliography_field_preserves_field_and_text(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from zipfile import ZipFile
    doc = Document()
    doc.add_paragraph('Final reference remains unchanged')
    paragraph = doc.add_paragraph()
    ending = OxmlElement('w:fldChar'); ending.set(qn('w:fldCharType'), 'end')
    paragraph.add_run()._r.append(ending)
    original, target = tmp_path/'source.docx', tmp_path/'revised.docx'
    doc.save(original)
    payload = {'source_sha256': hashlib.sha256(original.read_bytes()).hexdigest(),
               'compact_trailing_field_paragraph': True}
    build(original, payload, target)
    assert inspect(original)['paragraphs'] == inspect(target)['paragraphs']
    with ZipFile(target) as archive:
        xml = archive.read('word/document.xml')
        assert b'w:fldCharType="end"' in xml and b'w:line="20"' in xml


@pytest.mark.parametrize("changed", ["artifact", "existing_delivery"])
def test_publication_rejects_stale_artifacts_and_preserves_user_edits(tmp_path, monkeypatch, changed):
    from paper_analysis.teacher import hf_qc_delivery as module
    monkeypatch.setattr(module, "check", lambda _: [])
    monkeypatch.setattr(module, "git_sha", lambda _: "committed")
    monkeypatch.setattr(module.subprocess, "check_output", lambda *a, **kw: "")
    work, delivery = tmp_path / "work", tmp_path / "teacher"
    stage = work / "staging"
    stage.mkdir(parents=True)
    delivery.mkdir()
    (delivery / "manual.txt").write_text("user work")
    report = stage / "论文数据分析结果报告.md"
    report.write_text("verified report")
    digest = hashlib.sha256(report.read_bytes()).hexdigest()
    config = {"run_root": str(tmp_path / "run"), "work_root": str(work), "delivery_root": str(delivery), "archive_path": str(tmp_path / "archive")}
    prepared = {"code_sha": "committed", "config_sha256": hashlib.sha256(json.dumps(config, sort_keys=True).encode()).hexdigest(),
        "old_inventory": module.inventory(delivery), "artifacts": {report.name: digest}}
    (work / "preparation.json").write_text(json.dumps(prepared))
    (work / "artifact_verification.json").write_text(json.dumps({"files": {report.name: {"sha256": digest, "content_verified": True}}}))
    if changed == "artifact": report.write_text("changed")
    else: (delivery / "manual.txt").write_text("new user changes")
    current = (delivery / "manual.txt").read_bytes()
    with pytest.raises(StageBlockedError, match="stale|Teacher delivery changed"):
        module.publish(config, tmp_path)
    assert (delivery / "manual.txt").read_bytes() == current
    assert stage.exists()
