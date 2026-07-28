from __future__ import annotations

from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def write_markdown_report(
    path: str | Path,
    *,
    title: str,
    paragraphs: Iterable[str],
    tables: Iterable[tuple[str, pd.DataFrame]] = (),
) -> Path:
    """Write a readable report while keeping complete tables as separate artifacts."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    lines = [f"# {title}", ""]
    for paragraph in paragraphs:
        lines.extend([str(paragraph), ""])
    for heading, frame in tables:
        lines.extend([f"## {heading}", ""])
        preview = frame.head(15)
        if preview.empty:
            lines.extend(["No rows.", ""])
            continue
        if len(preview.columns) > 8:
            preview = preview.iloc[:, :8].copy()
            lines.extend([
                "下表仅展示前 15 行和前 8 列；完整机器可读表随结果单独交付。",
                "",
            ])
        lines.append("```csv")
        lines.append(preview.to_csv(index=False, lineterminator="\n").rstrip())
        lines.extend(["```", ""])
    target.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return target


def write_docx_report(
    path: str | Path,
    *,
    title: str,
    paragraphs: Iterable[str],
    tables: Iterable[tuple[str, pd.DataFrame]] = (),
) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Title", 23, "000000", 0, 10),
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("Teacher-priority analysis report")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("Generated from machine-readable analysis outputs")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("777777")
    document.add_heading(title, 0)
    for paragraph in paragraphs:
        document.add_paragraph(str(paragraph))
    for heading, frame in tables:
        document.add_heading(heading, level=1)
        preview = frame.head(15)
        if preview.empty:
            document.add_paragraph("No rows.")
            continue
        if len(preview.columns) > 8:
            preview = preview.iloc[:, :8].copy()
            document.add_paragraph(
                "The table preview shows the first 8 columns; the complete "
                "machine-readable table is delivered separately."
            )
        table = document.add_table(rows=1, cols=len(preview.columns))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        column_width = Inches(6.5 / max(len(preview.columns), 1))
        for index, column in enumerate(preview.columns):
            table.rows[0].cells[index].text = str(column)
            cell = table.rows[0].cells[index]
            cell.width = column_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F2F4F7")
            cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for cell_run in paragraph.runs:
                    cell_run.bold = True
                    cell_run.font.size = Pt(8)
        for _, row in preview.iterrows():
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = "" if pd.isna(value) else str(value)
                cells[index].width = column_width
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cells[index].paragraphs:
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if pd.api.types.is_number(value)
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    for cell_run in paragraph.runs:
                        cell_run.font.size = Pt(8)
    document.save(target)
    return target


def save_condition_plot(
    frame: pd.DataFrame,
    *,
    outcome: str,
    path: str | Path,
    title: str,
) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(7.2, 4.6), constrained_layout=True)
    if not frame.empty and outcome in frame:
        summary = (
            frame.groupby(["WWR", "Complexity"], dropna=False)[outcome]
            .agg(["mean", "sem"])
            .reset_index()
        )
        for complexity, sub in summary.groupby("Complexity", dropna=False):
            ax.errorbar(
                sub["WWR"].astype(str), sub["mean"], yerr=1.96 * sub["sem"],
                marker="o", capsize=3, label=str(complexity),
            )
        ax.legend(title="Complexity", frameon=False)
    else:
        ax.text(.5, .5, "No estimable data", ha="center", va="center")
    ax.set_title(title)
    ax.set_xlabel("WWR (categorical)")
    ax.set_ylabel(outcome)
    fig.savefig(target, dpi=300)
    plt.close(fig)
    return target


def save_histogram(
    values: pd.Series,
    *,
    path: str | Path,
    title: str,
    xlabel: str,
    thresholds: Iterable[float] = (),
) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(7.2, 4.6), constrained_layout=True)
    numeric = pd.to_numeric(values, errors="coerce").dropna()
    if numeric.empty:
        ax.text(.5, .5, "No estimable data", ha="center", va="center")
    else:
        ax.hist(numeric, bins=min(20, max(5, int(numeric.nunique()))), color="#4C78A8")
    for threshold in thresholds:
        ax.axvline(float(threshold), linestyle="--", linewidth=1)
    ax.set_title(title)
    ax.set_xlabel(xlabel)
    ax.set_ylabel("Count")
    fig.savefig(target, dpi=300)
    plt.close(fig)
    return target
