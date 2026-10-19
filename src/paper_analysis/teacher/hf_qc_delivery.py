"""Focused, transactional publication of an audited No-HF revision package.

Scientific content is provided as a local, hash-bound writing payload. This
publisher validates its numeric claims and preserves the source manuscript.
"""
from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

from .fresh import dump, git_sha
from .state import StageBlockedError, file_sha256
from .reproducibility import verify_seal, verify_complete_bootstrap_repeat
from .hf_qc_revision import read_json
from .fresh_handoff import SHEETS, typed
from .request_handoff import archive_publish


def scientific_run(config):
    """Keep immutable scientific evidence separate from a later writing run."""
    source = Path(config.get("scientific_run_root", config["run_root"]))
    output = Path(config["run_root"])
    if source.resolve() != output.resolve() and (
        source.resolve().is_relative_to(output.resolve())
        or output.resolve().is_relative_to(source.resolve())
    ):
        raise StageBlockedError("Scientific and writing runs must be disjoint")
    return source


def scientific_pointer(pointer):
    """Analysis consumers must keep resolving real tables, not a writing folder."""
    return dict(pointer, run_root=pointer["scientific_run_root"],
                delivery_run_root=pointer["run_root"])


def validate_numeric_claims(payload, source_by_key):
    rows = []
    for item in payload.get("numeric_claims", []):
        path = Path(source_by_key[item["source"]]["source_path"])
        if path.suffix.lower() == ".csv":
            table = pd.read_csv(path, encoding="utf-8-sig", float_precision="round_trip")
        else:
            table = pd.read_excel(path, sheet_name=item.get("sheet", 0))
        for field, value in item.get("filters", {}).items():
            if isinstance(value, (int, float)):
                table = table.loc[pd.to_numeric(table[field], errors="coerce").eq(value)]
            else:
                table = table.loc[table[field].astype(str).eq(str(value))]
        if len(table) != 1:
            raise StageBlockedError("Numeric claim does not identify one source row: " + str(item))
        actual = table.iloc[0][item["field"]]
        if not np.isfinite(float(actual)) or not np.isclose(float(actual), float(item["expected"]), rtol=1e-10, atol=1e-12):
            raise StageBlockedError("Numeric claim differs from its source: " + str(item))
        rows.append({"位置": item["location"], "Source ID": source_by_key[item["source"]]["Source ID"],
            "字段": item["field"], "筛选条件": json.dumps(item.get("filters", {}), ensure_ascii=False),
            "文稿引用原值": item["expected"], "源表原值": actual, "核验": "一致",
            "平铺文件": source_by_key[item["source"]]["flat_name"]})
    if not rows:
        raise StageBlockedError("No source-backed numerical checks in writing payload")
    return pd.DataFrame(rows)


def check(config):
    run = scientific_run(config)
    work = Path(config["work_root"]).resolve()
    delivery = Path(config["delivery_root"]).resolve()
    archive = Path(config["archive_path"]).resolve()
    read_json(Path(config["outputs_root"]) / "realdata_run_summary.json")
    for a, b in [(work, delivery), (work, archive), (delivery, archive)]:
        if a == b or a.is_relative_to(b) or b.is_relative_to(a):
            raise StageBlockedError("Work, delivery and external archive must be disjoint")
    manifest = read_json(run / "run_manifest.json")
    if manifest["status"] != "analysis_complete_pending_delivery":
        raise StageBlockedError("Revised scientific run has not completed")
    if not read_json(run / "methodological_gate.json")["adopt_NoHF"]:
        raise StageBlockedError("Formal No-HF gate did not pass")
    proof = read_json(run / "complete_verification/verification.json")
    if proof.get("status") != "passed" or {k for k, v in proof["scopes"].items() if v["status"] == "passed"} != {"eye", "eeg", "joint"}:
        raise StageBlockedError("Eye, EEG and joint verification must all pass")
    verify_seal(proof["evidence"])
    verify_seal(proof["eye_reuse"])
    for item in proof["checks"]:
        if not item["passed"]: raise StageBlockedError("Failed numerical comparison")
        verify_seal(item["inputs"])
    verify_complete_bootstrap_repeat(run / "bootstrap_repeat/repeat_execution.json")
    if file_sha256(config["writing_payload"]) != config["writing_payload_sha256"]:
        raise StageBlockedError("Writing payload changed")
    payload = read_json(config["writing_payload"])
    if file_sha256(config["manuscript_file"]) != payload["manuscript"]["source_sha256"]:
        raise StageBlockedError("Original manuscript changed")
    records = []
    for i, item in enumerate(payload["sources"], 1):
        path = Path(item["path"])
        if file_sha256(path) != item["sha256"]:
            raise StageBlockedError("Source changed: " + str(path))
        records.append({"Source ID": f"SRC{i:04d}", "key": item["key"], "source_path": str(path),
            "flat_name": f"SRC{i:04d}__{item['stage']}__{path.name}", "stage": item["stage"],
            "bytes": path.stat().st_size, "SHA256": item["sha256"]})
    if len({r["key"] for r in records}) != len(records):
        raise StageBlockedError("Duplicate source key")
    validate_numeric_claims(payload, {r["key"]: r for r in records})
    return records


def build_docx(spec, target, sources):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    for border in list(doc.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    doc.styles["Normal"].font.size = Pt(10.5)
    if "title_size_pt" in spec:
        doc.styles["Title"].font.size = Pt(spec["title_size_pt"])
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.add_paragraph(spec["title"], "Title")
    for block in spec["blocks"]:
        if "heading" in block:
            doc.add_heading(block["heading"], level=block.get("level", 1))
        elif "paragraph" in block:
            doc.add_paragraph(block["paragraph"])
        elif "table" in block:
            rows = block["table"]
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = False
            weights = [max(8, min(28, max(len(str(row[i])) for row in rows))) for i in range(len(rows[0]))]
            borders = OxmlElement("w:tblBorders")
            for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                border = OxmlElement("w:"+edge)
                for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]: border.set(qn("w:"+key), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            for c, value in zip(table.rows[0].cells, rows[0]): c.text = str(value)
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)
            for values in rows[1:]:
                for c, value in zip(table.add_row().cells, values): c.text = str(value)
            for ri, row in enumerate(table.rows):
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for ci, cell in enumerate(row.cells):
                    cell.width = Inches(6.9 * weights[ci] / sum(weights))
                    props = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for edge in ["top", "left", "bottom", "right"]:
                        margin = OxmlElement("w:"+edge); margin.set(qn("w:w"), "90"); margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    props.append(margins)
                    if ri == 0:
                        fill = OxmlElement("w:shd"); fill.set(qn("w:fill"), "E7E6E6"); props.append(fill)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(3)
                        for r in paragraph.runs:
                            r.font.size = Pt(9)
                            r.font.bold = ri == 0
        if block.get("sources"):
            note = doc.add_paragraph("来源：" + "；".join(sources[k]["Source ID"] for k in block["sources"]) + "（文件与筛选条件见交接索引）")
            note.paragraph_format.space_before = Pt(0)
            note.paragraph_format.space_after = Pt(3)
            for run in note.runs: run.font.size = Pt(9)
    section.footer.paragraphs[0].alignment = 2
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    section.footer.paragraphs[0]._p.append(field)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    from .document_fonts import apply_document_fonts
    apply_document_fonts(target)


def prepare(config, repo):
    if subprocess.check_output(["git", "status", "--porcelain"], cwd=repo, text=True).strip():
        raise StageBlockedError("Prepare from clean committed code")
    records = check(config)
    run, work = Path(config["run_root"]), Path(config["work_root"])
    run.mkdir(parents=True, exist_ok=True)
    stage = work / "staging"
    stage.mkdir(parents=True, exist_ok=False)
    flat = stage / "filesource_flat"
    flat.mkdir()
    payload = read_json(config["writing_payload"])
    sources = {r["key"]: r for r in records}
    for item in records:
        shutil.copyfile(item["source_path"], flat / item["flat_name"])
        if file_sha256(flat / item["flat_name"]) != item["SHA256"]:
            raise StageBlockedError("Source copy differs")
    dump(work / "source_manifest.json", records)
    dump(flat / "source_manifest.json", records)
    numbers = validate_numeric_claims(payload, sources)
    numbers.to_csv(work / "numeric_claim_verification.csv", index=False, encoding="utf-8-sig")
    report = payload["report_md"]
    for key, value in sources.items(): report = report.replace("{{" + key + "}}", value["Source ID"])
    if "{{" in report:
        raise StageBlockedError("Unresolved report Source ID")
    (run / "论文数据分析结果报告.md").write_text(report, encoding="utf-8")
    shutil.copyfile(run / "论文数据分析结果报告.md", stage / "论文数据分析结果报告.md")
    # Mark authoring immediately before the first artifact operation.
    subprocess.run([config["node"], str(Path(config["documents_skill"]) / "container_tools/mark_artifact_operation_started.mjs"),
        "--operation-kind", "create", "--expected-output-count", str(len(payload["documents"])+1), "--output-format", "docx"], check=True)
    for spec in payload["documents"]:
        build_docx(spec, stage / spec["name"], sources)
    from .manuscript_document import build
    changes = build(config["manuscript_file"], payload["manuscript"], stage / "manuscript_REVISED_NO_HF_QC.docx")
    dump(work / "manuscript_changes.json", changes)
    pd.DataFrame(changes).to_csv(run / "manuscript_changes.csv", index=False, encoding="utf-8-sig")
    mapping = []
    for item in payload["paragraph_map"]:
        for key in item["sources"]:
            r = sources[key]
            mapping.append({"文档": item["artifact"], "位置": item["location"], "内容摘要": item["summary"],
                "Source ID": r["Source ID"], "工作表或字段": item.get("fields", ""), "筛选条件": item.get("filters", ""),
                "核验状态": item.get("status", "源值核对通过"), "平铺文件": r["flat_name"]})
    if not mapping:
        raise StageBlockedError("Paragraph-to-source mapping is required")
    mapping = pd.DataFrame(mapping)
    mapping.to_csv(work / "paragraph_source_map.csv", index=False, encoding="utf-8-sig")
    table_data = [pd.DataFrame(payload["overview"]), mapping, numbers, pd.DataFrame(records).drop(columns="key"),
        pd.DataFrame(payload["sample_rows"]), pd.DataFrame(payload["verification_rows"]),
        pd.DataFrame(payload["differences"]), pd.DataFrame(payload["external_sources"]), pd.DataFrame(payload["terms"])]
    tables = {"tables": [typed(name, frame) for name, frame in zip(SHEETS, table_data)]}
    dump(work / "index_payload.json", tables)
    helpers = work / "helpers"
    helpers.mkdir()
    shutil.copyfile(repo / "scripts/build_fresh_handoff_index.mjs", helpers / "index.mjs")
    packages = Path(config["node_modules"])
    subprocess.run(["cmd", "/c", "mklink", "/J", str(helpers / "node_modules"), str(packages)], check=True, capture_output=True)
    subprocess.run([config["node"], str(Path(config["spreadsheets_skill"]) / "container_tools/mark_artifact_operation_started.mjs"),
        "--operation-kind", "create", "--expected-output-count", "2", "--output-format", "xlsx"], check=True)
    subprocess.run([config["node"], str(helpers / "index.mjs"), str(work / "index_payload.json"),
        str(stage / "数据来源交接索引.xlsx"), str(work / "xlsx_previews")], check=True)
    from .xlsx_hyperlinks import repair
    repaired = work / "index_links_finalized.xlsx"
    repair(stage / "数据来源交接索引.xlsx", repaired)
    shutil.copyfile(repaired, stage / "数据来源交接索引.xlsx")
    formal_payload = dict(tables, linkPrefix=Path(config["delivery_root"]).as_posix()+"/filesource_flat/")
    dump(work / "formal_index_payload.json", formal_payload)
    subprocess.run([config["node"], str(helpers / "index.mjs"), str(work / "formal_index_payload.json"),
        str(work / "结果文件总索引.xlsx"), str(work / "xlsx_previews/formal")], check=True)
    repair(work / "结果文件总索引.xlsx", work / "formal_index_finalized.xlsx")
    shutil.copyfile(work / "formal_index_finalized.xlsx", work / "结果文件总索引.xlsx")
    artifact_paths = [p for p in stage.iterdir() if p.is_file()]
    dump(work / "preparation.json", {"code_sha": git_sha(repo), "config_sha256": hashlib.sha256(json.dumps(config, sort_keys=True).encode()).hexdigest(),
        "scientific_run_root": str(scientific_run(config)),
        "calculation_sha": read_json(scientific_run(config) / "run_manifest.json")["code_sha"], "old_inventory": inventory(Path(config["delivery_root"])),
        "formal_index_sha256": file_sha256(work / "结果文件总索引.xlsx"),
        "artifacts": {p.name: file_sha256(p) for p in artifact_paths}})
    return stage


def inventory(root):
    return {p.relative_to(root).as_posix(): file_sha256(p) for p in sorted(root.rglob("*")) if p.is_file()}


def formal_document_sources(delivery):
    documents = {p.name: p for p in delivery.glob("*.docx")}
    revised = documents.get("manuscript_REVISED_NO_HF_QC.docx")
    if revised:
        documents["论文全文正文_最新结果直接修改版.docx"] = revised
    return documents


def publish(config, repo):
    if subprocess.check_output(["git", "status", "--porcelain"], cwd=repo, text=True).strip():
        raise StageBlockedError("Publish from clean committed code")
    records = check(config)
    run, work = Path(config["run_root"]), Path(config["work_root"])
    stage = work / "staging"
    prepared = read_json(work / "preparation.json")
    retired = config.get("retired_formal_documents", [])
    for name in retired:
        if (Path(name).name != name or Path(name).suffix.lower() not in {".docx", ".csv"}
                or name in prepared["artifacts"]):
            raise StageBlockedError("Invalid retired formal document name")
    if prepared["code_sha"] != git_sha(repo):
        raise StageBlockedError("Publication code changed after preparation")
    if prepared["config_sha256"] != hashlib.sha256(json.dumps(config, sort_keys=True).encode()).hexdigest():
        raise StageBlockedError("Publication configuration changed")
    qa = read_json(work / "artifact_verification.json")
    for name, expected in prepared["artifacts"].items():
        item = qa["files"][name]
        if file_sha256(stage / name) != expected or item["sha256"] != expected or not item["content_verified"]:
            raise StageBlockedError("Artifact verification stale or incomplete: " + name)
        if name.endswith((".docx", ".xlsx")) and not item["visually_reviewed"]:
            raise StageBlockedError("Visual verification missing: " + name)
    for item in records:
        if file_sha256(stage / "filesource_flat" / item["flat_name"]) != item["SHA256"]:
            raise StageBlockedError("Staged source copy changed")
    index_sha = prepared.get("formal_index_sha256")
    if index_sha:
        item = qa["files"]["结果文件总索引.xlsx"]
        if file_sha256(work / "结果文件总索引.xlsx") != index_sha or item["sha256"] != index_sha or not item["content_verified"] or not item["visually_reviewed"]:
            raise StageBlockedError("Formal index verification stale or incomplete")
    destination, archive = Path(config["delivery_root"]), Path(config["archive_path"])
    if inventory(destination) != prepared["old_inventory"]:
        raise StageBlockedError("Teacher delivery changed after preparation; preserve and rebuild")
    files = inventory(stage)
    dump(stage / "filesource_flat/交付文件哈希.json", files)
    files = inventory(stage)
    archive_file = stage / config["zip_name"]
    with zipfile.ZipFile(archive_file, "w", zipfile.ZIP_DEFLATED) as z:
        for name in files: z.write(stage / name, name)
    extracted = work / "zip_extracted_verification"
    with zipfile.ZipFile(archive_file) as z:
        if z.testzip() is not None: raise StageBlockedError("ZIP CRC failure")
        z.extractall(extracted)
    if inventory(extracted) != files:
        raise StageBlockedError("Extracted ZIP inventory differs")
    # Offline workbook links must point into the actual extracted source folder.
    from openpyxl import load_workbook
    workbook = load_workbook(extracted / "数据来源交接索引.xlsx", data_only=False)
    if workbook.sheetnames != SHEETS:
        raise StageBlockedError("Nine-sheet workbook structure differs")
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('=HYPERLINK("filesource_flat/'):
                    name = cell.value.split('"')[1]
                    if not (extracted / name).is_file(): raise StageBlockedError("Offline workbook link broken")
    numeric_book = load_workbook(extracted / "数据来源交接索引.xlsx", data_only=True)
    expected_numbers = pd.read_csv(work / "numeric_claim_verification.csv")
    number_sheet = numeric_book["关键数字核对"]
    if number_sheet.max_row - 1 != len(expected_numbers):
        raise StageBlockedError("Workbook numeric-check row count differs")
    for index, record in enumerate(expected_numbers.itertuples(index=False, name=None), 2):
        for column in [5, 6]:
            if not np.isclose(float(number_sheet.cell(index, column).value), float(record[column-1]), rtol=1e-10, atol=1e-12):
                raise StageBlockedError("Workbook numeric-check value differs")
    files[archive_file.name] = file_sha256(archive_file)
    def validate(current, historical):
        if inventory(historical) != prepared["old_inventory"]: raise StageBlockedError("Old delivery archive differs")
        if inventory(current) != files: raise StageBlockedError("Published delivery differs")
    moved = archive_publish(stage, destination, archive, validate)
    dump(work / "archive_path_mapping.json", moved)
    pointer = {"status": "complete", "scope": "formal_no_hf_eeg_and_affected_joint; eye_main_verified_reuse; questionnaire_not_rerun",
        "run_root": str(run), "work_root": str(work), "analysis_git_sha": prepared["calculation_sha"],
        "scientific_run_root": str(scientific_run(config)),
        "publication_git_sha": git_sha(repo), "delivery_root": str(destination), "archive_root": str(archive),
        "zip": str(destination / archive_file.name), "zip_sha256": files[archive_file.name], "source_count": len(records),
        "analysis_config": str(Path(config["outputs_root"]) / "analysis_latest_eye_eeg.local.json"),
        "archived_input_resolution": str(work / "archive_path_mapping.json")}
    outputs = Path(config["outputs_root"])
    previous = read_json(outputs / "realdata_run_summary.json")
    dump(work / "previous_realdata_run_summary.json", previous)
    payload = read_json(config["writing_payload"])
    pointer_names = ["hf_qc_revision_latest.json", "eye_eeg_fresh_latest.json", "eeg_denominator_sensitivity_latest.json",
                     "eeg_request_handoff_latest.json", "eye_eeg_incremental_latest.json", "manuscript_delivery_latest.json", "results_revision_note_latest.json"]
    targets = [outputs / name for name in ["realdata_run_summary.json", "realdata_run_summary.csv", "论文数据分析结果报告.md",
                                          "README_当前有效结果.md", "README_老师本次EEG交付.md", "结果文件总索引.xlsx", "code_publication.json", "analysis_latest_eye_eeg.local.json", *pointer_names]]
    documents = formal_document_sources(destination)
    targets.extend(outputs / name for name in documents)
    targets.extend(outputs / name for name in retired)
    backups = {path: path.read_bytes() if path.exists() else None for path in targets}
    try:
        preserved = work / "previous_formal_documents"
        preserved.mkdir()
        for name in [*documents, *retired]:
            if (outputs / name).is_file(): shutil.copyfile(outputs / name, preserved / name)
        summary = {"status": "complete", "run_root": str(scientific_run(config)),
            "delivery_run_root": str(run), "analysis_git_sha": pointer["analysis_git_sha"],
            "publication_git_sha": pointer["publication_git_sha"], "scope_note": pointer["scope"],
            "samples": payload["scientific_samples"], "eye": payload["scientific_samples"]["eye"],
            "eeg": payload["scientific_samples"]["eeg"], "joint": payload["scientific_samples"]["joint"],
            "synchronized": payload["scientific_samples"]["joint"]["synchronized"],
            "questionnaire": {"status": "not_rerun", "original_53_person_chapter_source": "not_fully_verified"},
            "reproducibility": {"status": "passed", "proof": str(scientific_run(config) / "complete_verification/verification.json")},
            "scientific_run_root": str(scientific_run(config)),
            "teacher_delivery": pointer, "current_teacher_request": pointer, "latest_hf_qc_revision": pointer,
            "latest_manuscript_delivery": pointer, "previous_summary": str(work / "previous_realdata_run_summary.json")}
        dump(outputs / "realdata_run_summary.json", summary)
        pd.DataFrame([{"run_id": run.name, "status": "complete", "analysis_git_sha": pointer["analysis_git_sha"],
            "eye_participants": summary["eye"]["participants"], "eye_trials": summary["eye"]["trials"],
            "eeg_participants": summary["eeg"]["participants"], "eeg_trials": summary["eeg"]["trials"],
            "QC": "NoHF", "questionnaire": "not_rerun"}]).to_csv(outputs / "realdata_run_summary.csv", index=False, encoding="utf-8-sig")
        analysis_pointers = {"hf_qc_revision_latest.json", "eye_eeg_fresh_latest.json",
                             "eeg_denominator_sensitivity_latest.json", "eye_eeg_incremental_latest.json"}
        for name in pointer_names:
            dump(outputs / name, scientific_pointer(pointer) if name in analysis_pointers else pointer)
        for p in destination.iterdir():
            if p.is_file() and p.suffix.lower() in {".docx", ".md", ".xlsx"}: shutil.copyfile(p, run / p.name)
        shutil.copyfile(run / "论文数据分析结果报告.md", outputs / "论文数据分析结果报告.md")
        shutil.copyfile(work / "结果文件总索引.xlsx", outputs / "结果文件总索引.xlsx")
        for name, source in documents.items(): shutil.copyfile(source, outputs / name)
        for name in retired:
            if (outputs / name).is_file(): (outputs / name).unlink()
        dump(outputs / "code_publication.json", pointer)
        active_config = read_json(scientific_run(config) / "resolved_config.json")
        active_config["run_id"] = "eye_eeg_no_hf_next_run"
        dump(outputs / "analysis_latest_eye_eeg.local.json", active_config)
        (outputs / "README_当前有效结果.md").write_text("# 当前有效结果\n\n正式 EEG QC 已修订为 No-HF，相关 EEG 和联合结果以本轮为准。眼动独立主分析经核验复用，问卷独立结果未重跑。\n\n" + str(run) + "\n\n" + pointer["zip"] + "\n", encoding="utf-8")
        shutil.copyfile(outputs / "README_当前有效结果.md", outputs / "README_老师本次EEG交付.md")
    except BaseException:
        for path, data in backups.items():
            if data is None:
                if path.exists(): path.unlink()
            else: path.write_bytes(data)
        def validate_restored(current, _):
            if inventory(current) != prepared["old_inventory"]:
                raise StageBlockedError("Rollback differs")
        archive_publish(archive, destination, work / "failed_publication", validate_restored)
        raise
    dump(work / "publication.json", pointer)
    dump(run / "publication.json", pointer)
    return pointer
